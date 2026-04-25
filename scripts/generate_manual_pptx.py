#!/usr/bin/env python3
"""Generate user_manual.pptx from HTML manual content."""

import io, os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BASE   = os.path.join(os.path.dirname(__file__), '..')
ICON   = os.path.join(BASE, 'pymasking', 'web', 'static', 'favicon.png')
GIF_T  = os.path.join(BASE, 'doc', 'demo_text.gif')
GIF_F  = os.path.join(BASE, 'doc', 'demo_file.gif')
GIF_C  = os.path.join(BASE, 'doc', 'demo_clipboard.gif')
OUT    = os.path.join(BASE, 'user_manual.pptx')

# ── 色 ───────────────────────────────────────────────────────────────────
C_DARK   = RGBColor(0x7c, 0x3a, 0x00)   # 見出し濃茶
C_ORANGE = RGBColor(0xf4, 0xa2, 0x61)   # アクセントオレンジ
C_BG     = RGBColor(0xfd, 0xf6, 0xf0)   # スライド背景
C_WHITE  = RGBColor(0xff, 0xff, 0xff)
C_BODY   = RGBColor(0x33, 0x33, 0x33)
C_CODE   = RGBColor(0xa0, 0x44, 0x0a)
C_CODE_BG= RGBColor(0xff, 0xf3, 0xe8)
C_NOTE   = RGBColor(0x77, 0x77, 0x77)

# スライドサイズ: 16:9 ワイド
SW = Inches(13.33)
SH = Inches(7.5)

FONT_JP = "MS Gothic"

prs = Presentation()
prs.slide_width  = SW
prs.slide_height = SH

blank_layout = prs.slide_layouts[6]   # 完全ブランク

# ── ヘルパー ─────────────────────────────────────────────────────────────
def add_slide():
    slide = prs.slides.add_slide(blank_layout)
    # 背景色
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = C_BG
    return slide

def tf(shape):
    return shape.text_frame

def add_rect(slide, x, y, w, h, fill_rgb=None, line_rgb=None, line_w=Pt(0)):
    from pptx.util import Emu
    shape = slide.shapes.add_shape(1, x, y, w, h)  # MSO_SHAPE_TYPE.RECTANGLE=1
    shape.line.width = line_w
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if line_rgb:
        shape.line.color.rgb = line_rgb
    else:
        shape.line.fill.background()
    return shape

def add_textbox(slide, x, y, w, h, text, size=Pt(18), bold=False,
                color=C_BODY, align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(x, y, w, h)
    txb.word_wrap = wrap
    p = txb.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT_JP
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    return txb

def add_para(tf_obj, text, size=Pt(16), bold=False, color=C_BODY,
             indent=0, align=PP_ALIGN.LEFT, space_before=Pt(0)):
    from pptx.util import Pt as _Pt
    p = tf_obj.add_paragraph()
    p.alignment = align
    p.space_before = space_before
    if indent:
        p.level = indent
    run = p.add_run()
    run.text = text
    run.font.name = FONT_JP
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    return p

def header_bar(slide, title, icon_path=None):
    """濃茶ヘッダーバー"""
    bar = add_rect(slide, 0, 0, SW, Inches(1.1), fill_rgb=C_DARK)
    if icon_path and os.path.exists(icon_path):
        slide.shapes.add_picture(icon_path, Inches(0.25), Inches(0.18), Inches(0.72), Inches(0.72))
    add_textbox(slide, Inches(1.2), Inches(0.2), Inches(11), Inches(0.72),
                title, size=Pt(28), bold=True, color=C_WHITE)

def section_title(slide, text, y=Inches(1.3)):
    """オレンジ左ボーダー付き見出し"""
    add_rect(slide, Inches(0.5), y, Inches(0.07), Inches(0.46), fill_rgb=C_ORANGE)
    add_textbox(slide, Inches(0.7), y, Inches(12), Inches(0.46),
                text, size=Pt(22), bold=True, color=C_DARK)
    return y + Inches(0.65)

def step_badge(slide, num, x, y, size=Inches(0.45)):
    """丸バッジ（番号）"""
    from pptx.util import Pt as _Pt
    badge = slide.shapes.add_shape(9, x, y, size, size)  # 9=OVAL
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_ORANGE
    badge.line.fill.background()
    p = badge.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(num)
    run.font.name = FONT_JP
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = C_WHITE

def gif_first_frame(gif_path):
    """GIF の最初のフレームを PNG バイト列で返す"""
    img = Image.open(gif_path)
    img.seek(0)
    buf = io.BytesIO()
    img.convert('RGB').save(buf, format='PNG')
    buf.seek(0)
    return buf

# ═══════════════════════════════════════════════════════════════════════════
# スライド 1: タイトル
# ═══════════════════════════════════════════════════════════════════════════
slide = add_slide()
add_rect(slide, 0, 0, SW, SH, fill_rgb=C_DARK)

if os.path.exists(ICON):
    slide.shapes.add_picture(ICON, Inches(5.4), Inches(1.4), Inches(2.5), Inches(2.5))

add_textbox(slide, Inches(1), Inches(4.2), Inches(11.3), Inches(1.0),
            "データマスキングツール マニュアル",
            size=Pt(34), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1), Inches(5.3), Inches(11.3), Inches(0.6),
            "pymasking",
            size=Pt(20), color=C_ORANGE, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════════════════════
# スライド 2: 目次
# ═══════════════════════════════════════════════════════════════════════════
slide = add_slide()
header_bar(slide, "目次", ICON)

items = [
    ("1", "インストール方法"),
    ("2", "起動方法"),
    ("3", "終了方法"),
    ("4", "データマスキング 3 種"),
]
y = Inches(1.55)
for num, label in items:
    step_badge(slide, num, Inches(0.9), y + Inches(0.01), size=Inches(0.42))
    add_textbox(slide, Inches(1.55), y, Inches(10), Inches(0.45),
                label, size=Pt(20), color=C_BODY)
    y += Inches(0.75)

# ═══════════════════════════════════════════════════════════════════════════
# スライド 3: インストール方法
# ═══════════════════════════════════════════════════════════════════════════
slide = add_slide()
header_bar(slide, "インストール方法", ICON)
y = section_title(slide, "インストール方法")

steps = [
    ("1", ["pymasking.7z をダウンロードする"]),
    ("2", [
        "ダウンロードした pymasking.7z を",
        "Windows 11 の標準の解凍方法で展開する",
        "展開先：%LOCALAPPDATA%\\pymasking",
    ]),
    ("3", [
        "pymasking のショートカットをデスクトップにコピーする",
        "%LOCALAPPDATA%\\pymasking\\pymasking.lnk",
        "　→ デスクトップにコピーする",
    ]),
]

for num, lines in steps:
    step_badge(slide, num, Inches(0.5), y, size=Inches(0.45))
    x_text = Inches(1.15)
    for i, line in enumerate(lines):
        is_code = line.startswith('%') or line.startswith('　→')
        clr = C_CODE if is_code else C_BODY
        sz  = Pt(14) if is_code else Pt(16)
        add_textbox(slide, x_text, y + Inches(i * 0.36), Inches(11.5), Inches(0.38),
                    line, size=sz, color=clr)
    y += Inches(len(lines) * 0.36 + 0.3)

# ═══════════════════════════════════════════════════════════════════════════
# スライド 4: 起動方法
# ═══════════════════════════════════════════════════════════════════════════
slide = add_slide()
header_bar(slide, "起動方法", ICON)
y = section_title(slide, "起動方法")

boot_steps = [
    ("1", "pymasking.lnk をダブルクリックする"),
    ("2", "コマンドプロンプト（黒い画面）が起動し、自動的にブラウザが起動する"),
]
for num, text in boot_steps:
    step_badge(slide, num, Inches(0.5), y, size=Inches(0.45))
    add_textbox(slide, Inches(1.15), y, Inches(11.5), Inches(0.48),
                text, size=Pt(18), color=C_BODY)
    y += Inches(0.8)

# ═══════════════════════════════════════════════════════════════════════════
# スライド 5: 終了方法
# ═══════════════════════════════════════════════════════════════════════════
slide = add_slide()
header_bar(slide, "終了方法", ICON)
y = section_title(slide, "終了方法")

stop_steps = [
    ("1", "ブラウザを閉じる", None),
    ("2", "コマンドプロンプト（黒い画面）は 15 秒後に自動的に終了する",
          "終了しない場合は Ctrl + C を連打して終了させる"),
]
for num, text, note in stop_steps:
    step_badge(slide, num, Inches(0.5), y, size=Inches(0.45))
    add_textbox(slide, Inches(1.15), y, Inches(11.5), Inches(0.48),
                text, size=Pt(18), color=C_BODY)
    if note:
        add_textbox(slide, Inches(1.35), y + Inches(0.5), Inches(11), Inches(0.38),
                    "※ " + note, size=Pt(14), color=C_NOTE)
        y += Inches(0.4)
    y += Inches(0.8)

# ═══════════════════════════════════════════════════════════════════════════
# スライド 6〜8: データマスキング 3 種（各 GIF 1 枚ずつ）
# ═══════════════════════════════════════════════════════════════════════════
demos = [
    ("テキストマスキング",      GIF_T),
    ("ファイルマスキング",      GIF_F),
    ("クリップボードマスキング", GIF_C),
]
for idx, (title, gif_path) in enumerate(demos, 1):
    slide = add_slide()
    header_bar(slide, f"データマスキング 3 種  [{idx}/3]", ICON)
    y = section_title(slide, f"{idx}. {title}")

    if os.path.exists(gif_path):
        img_buf = gif_first_frame(gif_path)
        # 画像サイズを取得してアスペクト比を保持
        img = Image.open(gif_path)
        iw, ih = img.size
        max_w = Inches(11.8)
        max_h = Inches(5.2)
        ratio = min(max_w / iw, max_h / ih)
        disp_w = int(iw * ratio)
        disp_h = int(ih * ratio)
        left = (SW - disp_w) // 2
        slide.shapes.add_picture(img_buf, left, y, disp_w, disp_h)

# ═══════════════════════════════════════════════════════════════════════════
# 保存
# ═══════════════════════════════════════════════════════════════════════════
prs.save(OUT)
print(f"Saved: {OUT}")
