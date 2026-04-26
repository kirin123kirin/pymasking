"""user_manual.docx を生成するスクリプト。"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

DOC_DIR = Path(__file__).parent.parent / "doc"
ORANGE       = RGBColor(0xE6, 0x51, 0x00)   # deep orange
ORANGE_LIGHT = RGBColor(0xFF, 0xA0, 0x40)   # light orange
ORANGE_BG    = RGBColor(0xFF, 0xF3, 0xE0)   # very light orange
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
DARK         = RGBColor(0x33, 0x20, 0x00)


def set_cell_bg(cell, rgb: RGBColor):
    """表セルの背景色を設定する。"""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")
    tcPr.append(shd)


def add_heading(doc: Document, text: str, level: int = 1):
    """オレンジ系スタイルの見出しを追加する。"""
    if level == 1:
        tbl = doc.add_table(rows=1, cols=1)
        cell = tbl.rows[0].cells[0]
        set_cell_bg(cell, ORANGE)
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = cp.add_run(f"  {text}")
        r.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = WHITE
        for row in tbl.rows:
            for c in row.cells:
                tcPr = c._tc.get_or_add_tcPr()
                tcBorders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    border = OxmlElement(f"w:{side}")
                    border.set(qn("w:val"), "none")
                    tcBorders.append(border)
                tcPr.append(tcBorders)
        doc.add_paragraph()
        return tbl
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = ORANGE
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), f"{ORANGE_LIGHT[0]:02X}{ORANGE_LIGHT[1]:02X}{ORANGE_LIGHT[2]:02X}")
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p


def add_step(doc: Document, step: str, desc: str):
    """番号付き手順行を追加する。"""
    p = doc.add_paragraph(style="List Number")
    run_step = p.add_run(step)
    run_step.bold = True
    run_step.font.color.rgb = ORANGE
    if desc:
        run_desc = p.add_run(f"  {desc}")
        run_desc.font.color.rgb = DARK


def add_note(doc: Document, text: str):
    """注記ボックスを追加する。"""
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, ORANGE_BG)
    cp = cell.paragraphs[0]
    r = cp.add_run(f"ℹ  {text}")
    r.font.size = Pt(10)
    r.font.color.rgb = DARK
    doc.add_paragraph()


def build():
    doc = Document()

    # ページ余白
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # デフォルトフォント
    doc.styles["Normal"].font.name = "メイリオ"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].font.color.rgb = DARK

    # ─── タイトル ───────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("pymasking  ユーザーマニュアル")
    tr.bold = True
    tr.font.size = Pt(28)
    tr.font.color.rgb = ORANGE
    doc.add_paragraph()

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("個人情報マスキングツール")
    sr.font.size = Pt(13)
    sr.font.color.rgb = ORANGE_LIGHT
    doc.add_paragraph()

    # ─── 起動方法 ────────────────────────────────────────
    add_heading(doc, "起動方法", level=1)

    doc.add_paragraph("デスクトップの masking.bat をダブルクリックして起動します。")
    doc.add_paragraph("またはコマンドプロンプトで以下を実行します。")

    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, RGBColor(0x33, 0x20, 0x00))
    cp = cell.paragraphs[0]
    r = cp.add_run("masking.bat")
    r.font.name = "Consolas"
    r.font.size = Pt(13)
    r.font.color.rgb = ORANGE_LIGHT
    doc.add_paragraph()

    add_note(doc, "ブラウザが自動的に起動し、http://127.0.0.1:55963 が開きます。")

    # ─── 終了方法 ────────────────────────────────────────
    add_heading(doc, "終了方法", level=1)

    add_step(doc, "方法①", "ブラウザのタブまたはウィンドウを閉じる（サーバーが自動終了します）")
    add_step(doc, "方法②", "コマンドプロンプトで Ctrl + C を押す")
    doc.add_paragraph()

    # ─── 使用方法 ────────────────────────────────────────
    add_heading(doc, "使用方法", level=1)
    doc.add_paragraph()

    # ── テキスト ──
    add_heading(doc, "テキスト", level=2)
    doc.add_paragraph()

    add_step(doc, "①", "「テキスト」タブを選択し、テキストエリアに文章を貼り付ける")
    add_step(doc, "②", "「マスキング方式」を選択し、必要に応じて「マスキング対象」をクリックして種別を絞り込む")
    add_step(doc, "③", "「マスキング実行」をクリックする")
    add_step(doc, "④", "結果欄に変換後テキストが表示される")
    add_step(doc, "⑤", "「結果をコピー」でクリップボードにコピーする")
    doc.add_paragraph()

    gif_text = DOC_DIR / "demo_text.gif"
    if gif_text.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(gif_text), width=Inches(5.5))
    doc.add_paragraph()

    # ── ファイル ──
    add_heading(doc, "ファイル", level=2)
    doc.add_paragraph()

    add_step(doc, "①", "「ファイル」タブを選択する")
    add_step(doc, "②", "ファイルをドラッグ＆ドロップ、またはクリックして選択する")
    add_step(doc, "③", "「マスキング方式」を選択し、必要に応じて「マスキング対象」をクリックして種別を絞り込む")
    add_step(doc, "④", "「マスキング実行」をクリックする（ボタンが赤く点滅）")
    add_step(doc, "⑤", "ボタンが「ダウンロード」に変わったらクリックして保存する")
    add_note(doc, "出力ファイル名は元のファイル名に「_変換後」が付加されます（例: report_変換後.docx）")

    gif_file = DOC_DIR / "demo_file.gif"
    if gif_file.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(gif_file), width=Inches(5.5))

    out = DOC_DIR / "user_manual.docx"
    doc.save(str(out))
    print(f"saved: {out}")


if __name__ == "__main__":
    build()
