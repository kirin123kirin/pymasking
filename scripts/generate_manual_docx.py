#!/usr/bin/env python3
"""Generate user_manual.docx from HTML manual content."""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE  = os.path.join(os.path.dirname(__file__), '..')
ICON  = os.path.join(BASE, 'pymasking', 'web', 'static', 'favicon.png')
GIF_T = os.path.join(BASE, 'doc', 'demo_text.gif')
GIF_F = os.path.join(BASE, 'doc', 'demo_file.gif')
GIF_C = os.path.join(BASE, 'doc', 'demo_clipboard.gif')
OUT   = os.path.join(BASE, 'user_manual.docx')

FONT_JP  = "MS Gothic"
C_DARK   = RGBColor(0x7c, 0x3a, 0x00)
C_ORANGE = RGBColor(0xf4, 0xa2, 0x61)
C_WHITE  = RGBColor(0xff, 0xff, 0xff)
C_BODY   = RGBColor(0x33, 0x33, 0x33)
C_CODE   = RGBColor(0xa0, 0x44, 0x0a)
C_NOTE   = RGBColor(0x77, 0x77, 0x77)

doc = Document()

# ── ページ余白 ────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ── スタイル調整 ──────────────────────────────────────────────────────────
style_normal = doc.styles['Normal']
style_normal.font.name = FONT_JP
style_normal.font.size = Pt(11)
style_normal.font.color.rgb = C_BODY
style_normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_JP)

# ── ヘルパー ──────────────────────────────────────────────────────────────
def set_font(run, size=Pt(11), bold=False, color=C_BODY, italic=False):
    run.font.name = FONT_JP
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_JP)

def para_space(para, before=Pt(0), after=Pt(4)):
    pf = para.paragraph_format
    pf.space_before = before
    pf.space_after  = after

def set_para_shading(para, fill_hex):
    """段落背景色を設定"""
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  fill_hex)
    pPr.append(shd)

def set_para_border_left(para, color_hex, size='36'):
    """段落左ボーダー"""
    pPr  = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'),   'single')
    left.set(qn('w:sz'),    size)
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), color_hex)
    pBdr.append(left)
    pPr.append(pBdr)

def add_heading(text, level=1):
    """オレンジ左ボーダー付き見出し"""
    p = doc.add_paragraph()
    para_space(p, before=Pt(14), after=Pt(6))
    set_para_border_left(p, 'F4A261', size='48' if level == 1 else '30')
    run = p.add_run(text)
    size = Pt(18) if level == 1 else Pt(14)
    set_font(run, size=size, bold=True, color=C_DARK)
    p.paragraph_format.left_indent = Inches(0.1)
    return p

def add_body(text, indent=0):
    p = doc.add_paragraph()
    para_space(p, before=Pt(2), after=Pt(2))
    run = p.add_run(text)
    set_font(run, size=Pt(11))
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    return p

def add_step(num, *lines, note=None):
    """番号付きステップ行"""
    p = doc.add_paragraph()
    para_space(p, before=Pt(5), after=Pt(2))
    # 番号
    r_num = p.add_run(f"  {num}.  ")
    set_font(r_num, size=Pt(12), bold=True, color=C_ORANGE)
    # 本文（最初の行）
    r_txt = p.add_run(lines[0])
    set_font(r_txt, size=Pt(11))
    # 追加行
    for line in lines[1:]:
        sub = doc.add_paragraph()
        para_space(sub, before=Pt(0), after=Pt(2))
        sub.paragraph_format.left_indent = Inches(0.45)
        r = sub.add_run(line)
        clr = C_CODE if (line.startswith('%') or line.startswith('　→')) else C_BODY
        sz  = Pt(10) if clr == C_CODE else Pt(11)
        set_font(r, size=sz, color=clr)
    if note:
        np = doc.add_paragraph()
        para_space(np, before=Pt(0), after=Pt(4))
        np.paragraph_format.left_indent = Inches(0.45)
        r = np.add_run("※ " + note)
        set_font(r, size=Pt(10), color=C_NOTE, italic=True)

def add_code_line(text, indent=0.45):
    p = doc.add_paragraph()
    para_space(p, before=Pt(1), after=Pt(1))
    set_para_shading(p, 'FFF3E8')
    p.paragraph_format.left_indent = Inches(indent)
    run = p.add_run(text)
    set_font(run, size=Pt(10), color=C_CODE)
    return p

def add_gif(path, width=Inches(5.5)):
    if os.path.exists(path):
        p = doc.add_paragraph()
        para_space(p, before=Pt(4), after=Pt(8))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=width)

def add_page_break():
    doc.add_page_break()

def title_para(text, size=Pt(28), color=C_WHITE, align=WD_ALIGN_PARAGRAPH.CENTER,
               bold=True, shade=None):
    p = doc.add_paragraph()
    p.alignment = align
    para_space(p, before=Pt(0), after=Pt(6))
    if shade:
        set_para_shading(p, shade)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return p

# ═══════════════════════════════════════════════════════════════════════════
# タイトルページ
# ═══════════════════════════════════════════════════════════════════════════
for _ in range(4):
    p = doc.add_paragraph()
    para_space(p)
    set_para_shading(p, '7C3A00')

# アイコン + タイトル
p_icon = doc.add_paragraph()
p_icon.alignment = WD_ALIGN_PARAGRAPH.CENTER
para_space(p_icon, before=Pt(0), after=Pt(0))
set_para_shading(p_icon, '7C3A00')
if os.path.exists(ICON):
    p_icon.add_run().add_picture(ICON, width=Inches(1.2))

title_para("データマスキングツール マニュアル", size=Pt(26), color=C_WHITE, shade='7C3A00')
title_para("pymasking", size=Pt(16), color=C_ORANGE, shade='7C3A00')

for _ in range(4):
    p = doc.add_paragraph()
    para_space(p)
    set_para_shading(p, '7C3A00')

add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
# 目次
# ═══════════════════════════════════════════════════════════════════════════
add_heading("目次")
for num, item in enumerate(["インストール方法", "起動方法", "終了方法", "データマスキング 3 種"], 1):
    p = doc.add_paragraph()
    para_space(p, before=Pt(4), after=Pt(4))
    p.paragraph_format.left_indent = Inches(0.3)
    r = p.add_run(f"{num}.  {item}")
    set_font(r, size=Pt(13), color=C_DARK)

add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
# インストール方法
# ═══════════════════════════════════════════════════════════════════════════
add_heading("１．インストール方法")

add_step(1, "pymasking.7z をダウンロードする")

add_step(2,
    "ダウンロードした pymasking.7z を",
    "Windows 11 の標準の解凍方法で展開する",
    "展開先に以下のパスを貼り付ける：",
)
add_code_line("%LOCALAPPDATA%\\pymasking")

add_step(3,
    "pymasking のショートカットをデスクトップにコピーする",
)
add_code_line("%LOCALAPPDATA%\\pymasking\\pymasking.lnk")
p = doc.add_paragraph()
para_space(p, before=Pt(0), after=Pt(2))
p.paragraph_format.left_indent = Inches(0.65)
r = p.add_run("　→ デスクトップにコピーする")
set_font(r, size=Pt(11))

add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
# 起動方法
# ═══════════════════════════════════════════════════════════════════════════
add_heading("２．起動方法")
add_step(1, "pymasking.lnk をダブルクリックする")
add_step(2, "コマンドプロンプト（黒い画面）が起動し、自動的にブラウザが起動する")

doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════════
# 終了方法
# ═══════════════════════════════════════════════════════════════════════════
add_heading("３．終了方法")
add_step(1, "ブラウザを閉じる")
add_step(2,
    "コマンドプロンプト（黒い画面）は 15 秒後に自動的に終了する",
    note="終了しない場合は Ctrl + C を連打して終了させる",
)

add_page_break()

# ═══════════════════════════════════════════════════════════════════════════
# データマスキング 3 種
# ═══════════════════════════════════════════════════════════════════════════
add_heading("４．データマスキング 3 種")

demos = [
    ("①  テキストマスキング",      GIF_T),
    ("②  ファイルマスキング",      GIF_F),
    ("③  クリップボードマスキング", GIF_C),
]
for label, gif in demos:
    add_heading(label, level=2)
    add_gif(gif, width=Inches(5.8))

# ═══════════════════════════════════════════════════════════════════════════
# 保存
# ═══════════════════════════════════════════════════════════════════════════
doc.save(OUT)
print(f"Saved: {OUT}")
