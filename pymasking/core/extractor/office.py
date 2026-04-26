"""Office ドキュメント（docx / xlsx / pptx）のマスキング処理。"""

import shutil
from pathlib import Path

from ..masker import mask_text, MaskMode
from . import make_output_path


def _mask_paragraph(para, mode: MaskMode, categories=None) -> None:
    raw = para.text
    if not raw.strip():
        return
    masked = mask_text(raw, mode, categories=categories)
    if masked == raw:
        return
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = masked
    else:
        para.add_run(masked)



def _remove_headers_footers(doc) -> None:
    """Remove all header and footer content from every section."""
    for section in doc.sections:
        for hf in [
            section.header, section.footer,
            section.even_page_header, section.even_page_footer,
            section.first_page_header, section.first_page_footer,
        ]:
            el = hf._element
            for child in list(el):
                el.remove(child)


def _remove_pptx_non_placeholders(prs) -> None:
    """Remove non-placeholder shapes (logos, copyright etc.) from slide masters and layouts."""
    for master in prs.slide_masters:
        for shape in list(master.shapes):
            if not shape.is_placeholder:
                shape._element.getparent().remove(shape._element)
        for layout in master.slide_layouts:
            for shape in list(layout.shapes):
                if not shape.is_placeholder:
                    shape._element.getparent().remove(shape._element)


def process_docx(src: Path, mode: MaskMode, categories=None, options=None) -> Path:
    from docx import Document

    options = options or {}
    out = make_output_path(src)
    shutil.copy2(src, out)
    doc = Document(out)

    if options.get("word_remove_headers"):
        _remove_headers_footers(doc)

    for para in doc.paragraphs:
        _mask_paragraph(para, mode, categories=categories)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _mask_paragraph(para, mode, categories=categories)

    doc.save(out)
    return out



def process_xlsx(src: Path, mode: MaskMode, categories=None, options=None) -> Path:
    from openpyxl import load_workbook

    out = make_output_path(src)
    shutil.copy2(src, out)
    wb = load_workbook(out)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    cell.value = mask_text(cell.value, mode, categories=categories)

    wb.save(out)
    return out



def process_pptx(src: Path, mode: MaskMode, categories=None, options=None) -> Path:
    from pptx import Presentation

    options = options or {}
    out = make_output_path(src)
    shutil.copy2(src, out)
    prs = Presentation(out)

    if options.get("pptx_remove_non_placeholders"):
        _remove_pptx_non_placeholders(prs)

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text:
                        run.text = mask_text(run.text, mode, categories=categories)

    prs.save(out)
    return out


def process_office(src: Path, mode: MaskMode, categories=None, options=None) -> Path:
    ext = src.suffix.lower()
    if ext == ".docx":
        return process_docx(src, mode, categories=categories, options=options)
    elif ext == ".xlsx":
        return process_xlsx(src, mode, categories=categories, options=options)
    elif ext == ".pptx":
        return process_pptx(src, mode, categories=categories, options=options)
    raise ValueError(f"未対応の Office 形式: {ext}")


